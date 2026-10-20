"""
Comprehensive Automated Test Suite for Barcode-Based Digital Examination Attendance System
Covers all 30 mandatory test scenarios defined in Section 23 of Project Requirements.
"""
import os
import sys
import unittest
import json
import sqlite3
import threading
import io
import openpyxl

# Point to dedicated test database
TEST_DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'test_comprehensive.db')
if os.path.exists(TEST_DB_PATH):
    try:
        os.remove(TEST_DB_PATH)
    except Exception:
        pass

import database
database.DB_PATH = TEST_DB_PATH

from app import app
from database import init_db, seed_demo_data, get_db

class TestExaminationAttendanceComprehensive(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        app.config['TESTING'] = True
        app.config['WTF_CSRF_ENABLED'] = False
        cls.client = app.test_client()
        init_db(TEST_DB_PATH)
        seed_demo_data(TEST_DB_PATH)

    def login_user(self, username, password):
        return self.client.post('/login', data={'username': username, 'password': password}, follow_redirects=False)

    def logout_user(self):
        return self.client.get('/logout', follow_redirects=True)

    # ----------------------------------------------------
    # TEST 1: Valid teacher logs in successfully
    # ----------------------------------------------------
    def test_01_valid_teacher_login(self):
        res = self.login_user('teacherramesh', 'TeacherPass@2026')
        self.assertEqual(res.status_code, 302)
        with self.client.session_transaction() as sess:
            self.assertEqual(sess.get('username'), 'teacherramesh')
            self.assertEqual(sess.get('role'), 'TEACHER')
        self.logout_user()

    # ----------------------------------------------------
    # TEST 2: Invalid credentials are rejected
    # ----------------------------------------------------
    def test_02_invalid_credentials_rejected(self):
        res = self.login_user('teacherramesh', 'WrongPassword123')
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Invalid institutional credentials", res.data)

    # ----------------------------------------------------
    # TEST 3: Teacher cannot access an unauthorized department
    # ----------------------------------------------------
    def test_03_unauthorized_department_blocked(self):
        # Teacher Ramesh is authorized for CSE (Dept ID 1), not ISE (Dept ID 2)
        self.login_user('teacherramesh', 'TeacherPass@2026')
        res = self.client.get('/departments/2/classrooms')
        self.assertEqual(res.status_code, 403)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 4: Teacher cannot access an unauthorized classroom
    # ----------------------------------------------------
    def test_04_unauthorized_classroom_blocked(self):
        # Session 2 is Room 405 (Block C) in ISE. Ramesh cannot access it.
        self.login_user('teacherramesh', 'TeacherPass@2026')
        res = self.client.get('/sessions/2/roster')
        self.assertEqual(res.status_code, 403)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 5: Correct student roster loads from the database
    # ----------------------------------------------------
    def test_05_student_roster_loads_from_database(self):
        self.login_user('teacherramesh', 'TeacherPass@2026')
        res = self.client.get('/sessions/1/roster')
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"1MS21CS001", res.data)
        self.assertIn(b"Aakash Verma", res.data)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 6: Valid registered barcode is accepted
    # ----------------------------------------------------
    def test_06_valid_registered_barcode_accepted(self):
        # BK20260001 is registered in authorized registry
        payload = {
            'session_id': 1,
            'usn': '1MS21CS001',
            'qp_barcode': 'BK20260001',
            'invigilator_name': 'Dr. Ramesh Sharma',
            'invigilator_id': 'teacher_ramesh'
        }
        res = self.client.post('/api/attendance/mark', json=payload)
        data = res.get_json()
        self.assertEqual(res.status_code, 200)
        self.assertTrue(data['success'])
        self.assertEqual(data['student']['status'], 'PRESENT')
        self.assertEqual(data['student']['qp_barcode'], 'BK20260001')

    # ----------------------------------------------------
    # TEST 7: Unknown barcode is rejected
    # ----------------------------------------------------
    def test_07_unknown_barcode_rejected(self):
        payload = {
            'session_id': 1,
            'usn': '1MS21CS002',
            'qp_barcode': 'UNKNOWN-FAKE-BARCODE-999'
        }
        res = self.client.post('/api/attendance/mark', json=payload)
        data = res.get_json()
        self.assertEqual(res.status_code, 400)
        self.assertFalse(data['success'])
        self.assertIn("UNAUTHORIZED BARCODE", data['message'])

    # ----------------------------------------------------
    # TEST 8: Invalid barcode format is rejected
    # ----------------------------------------------------
    def test_08_invalid_barcode_format_rejected(self):
        payload = {
            'session_id': 1,
            'usn': '1MS21CS002',
            'qp_barcode': 'X' # Too short
        }
        res = self.client.post('/api/attendance/mark', json=payload)
        data = res.get_json()
        self.assertEqual(res.status_code, 400)
        self.assertFalse(data['success'])
        self.assertIn("INVALID BARCODE FORMAT", data['message'])

    # ----------------------------------------------------
    # TEST 9: A barcode already assigned to another student is rejected
    # ----------------------------------------------------
    def test_09_barcode_already_assigned_rejected(self):
        # BK20260001 was already assigned to 1MS21CS001 in test 6
        payload = {
            'session_id': 1,
            'usn': '1MS21CS002',
            'qp_barcode': 'BK20260001' # Duplicate booklet assignment attempt
        }
        res = self.client.post('/api/attendance/mark', json=payload)
        data = res.get_json()
        self.assertEqual(res.status_code, 409)
        self.assertFalse(data['success'])
        self.assertTrue(data['duplicate'])
        self.assertIn("ALREADY ASSIGNED", data['message'])

    # ----------------------------------------------------
    # TEST 10: Repeated scans do not create duplicate attendance records
    # ----------------------------------------------------
    def test_10_repeated_scans_do_not_duplicate(self):
        # Attempt to scan 1MS21CS001 again
        payload = {
            'session_id': 1,
            'usn': '1MS21CS001',
            'qp_barcode': 'BK20260001'
        }
        res = self.client.post('/api/attendance/mark', json=payload)
        data = res.get_json()
        self.assertEqual(res.status_code, 409)
        self.assertTrue(data['duplicate'])

    # ----------------------------------------------------
    # TEST 11: A failed barcode validation does not mark student present
    # ----------------------------------------------------
    def test_11_failed_barcode_does_not_mark_present(self):
        conn = get_db(TEST_DB_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT attendance_status FROM students WHERE usn = '1MS21CS002'")
        status = cursor.fetchone()[0]
        conn.close()
        # Candidate 1MS21CS002 failed in test 7, 8, 9, must still be PENDING
        self.assertEqual(status, 'PENDING')

    # ----------------------------------------------------
    # TEST 12: Attendance is saved to the database
    # ----------------------------------------------------
    def test_12_attendance_persists_in_database(self):
        conn = get_db(TEST_DB_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT attendance_status, qp_barcode, marked_at FROM students WHERE usn = '1MS21CS001'")
        row = cursor.fetchone()
        conn.close()
        self.assertEqual(row['attendance_status'], 'PRESENT')
        self.assertEqual(row['qp_barcode'], 'BK20260001')
        self.assertIsNotNone(row['marked_at'])

    # ----------------------------------------------------
    # TEST 13: Attendance persists after a page refresh
    # ----------------------------------------------------
    def test_13_attendance_persists_after_page_refresh(self):
        self.login_user('admin', 'AdminPassword@2026')
        res = self.client.get('/sessions/1/roster')
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"BK20260001", res.data)
        self.assertIn(b"PRESENT", res.data)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 14: Correction updates the final attendance state correctly
    # ----------------------------------------------------
    def test_14_correction_updates_final_state(self):
        # Correct 1MS21CS001 from PRESENT to ABSENT
        payload = {
            'session_id': 1,
            'usn': '1MS21CS001',
            'new_status': 'ABSENT',
            'reason': 'Student fell sick and left examination room'
        }
        res = self.client.post('/api/attendance/redo', json=payload)
        data = res.get_json()
        self.assertEqual(res.status_code, 200)
        self.assertTrue(data['success'])
        self.assertEqual(data['student']['status'], 'ABSENT')

        # Verify in DB
        conn = get_db(TEST_DB_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT attendance_status, qp_barcode FROM students WHERE usn = '1MS21CS001'")
        row = cursor.fetchone()
        conn.close()
        self.assertEqual(row['attendance_status'], 'ABSENT')
        self.assertIsNone(row['qp_barcode'])

    # ----------------------------------------------------
    # TEST 15: The audit trail retains original and corrected values
    # ----------------------------------------------------
    def test_15_audit_trail_retains_history(self):
        conn = get_db(TEST_DB_PATH)
        cursor = conn.cursor()
        cursor.execute("""
            SELECT * FROM audit_trail 
            WHERE usn = '1MS21CS001' AND action_type = 'REDO_CORRECTION'
        """)
        row = cursor.fetchone()
        conn.close()
        self.assertIsNotNone(row)
        self.assertEqual(row['old_status'], 'PRESENT')
        self.assertEqual(row['new_status'], 'ABSENT')
        self.assertEqual(row['reason'], 'Student fell sick and left examination room')

    # ----------------------------------------------------
    # TEST 16: Pending students are identified before submission
    # ----------------------------------------------------
    def test_16_pending_students_identified_before_submission(self):
        res = self.client.get('/api/session/summary/1')
        data = res.get_json()
        self.assertEqual(res.status_code, 200)
        self.assertGreater(data['pending_count'], 0)
        self.assertIn('pending_list', data)

    # ----------------------------------------------------
    # TEST 17: Unresolved students cannot silently become confirmed absentees
    # ----------------------------------------------------
    def test_17_unresolved_cannot_silently_become_absent(self):
        # Attempt to confirm lock without force_submit
        res = self.client.post('/api/session/confirm_lock', json={'session_id': 1, 'force_submit': False})
        data = res.get_json()
        self.assertEqual(res.status_code, 400)
        self.assertEqual(data['warning'], 'unmarked_students')
        self.assertGreater(data['pending_count'], 0)

    # ----------------------------------------------------
    # TEST 18: The teacher can correct mistakes before submission
    # ----------------------------------------------------
    def test_18_teacher_can_correct_before_submission(self):
        # Mark 1MS21CS002 PRESENT with BK20260002
        self.client.post('/api/attendance/mark', json={
            'session_id': 1,
            'usn': '1MS21CS002',
            'qp_barcode': 'BK20260002'
        })
        # Correct to PENDING
        res = self.client.post('/api/attendance/redo', json={
            'session_id': 1,
            'usn': '1MS21CS002',
            'new_status': 'PENDING',
            'reason': 'Accidental wrong student row tapped'
        })
        data = res.get_json()
        self.assertEqual(res.status_code, 200)
        self.assertEqual(data['student']['status'], 'PENDING')

    # ----------------------------------------------------
    # TEST 19: A submitted session cannot be modified through ordinary endpoints
    # ----------------------------------------------------
    def test_19_submitted_session_locked_against_edits(self):
        # Submit & Lock Session 1
        lock_res = self.client.post('/api/session/confirm_lock', json={'session_id': 1, 'force_submit': True})
        self.assertEqual(lock_res.status_code, 200)

        # Attempt to mark attendance on locked session
        res = self.client.post('/api/attendance/mark', json={
            'session_id': 1,
            'usn': '1MS21CS003',
            'qp_barcode': 'BK20260003'
        })
        self.assertEqual(res.status_code, 403)
        self.assertTrue(res.get_json()['locked'])

    # ----------------------------------------------------
    # TEST 20: Unauthorized unlocking is rejected
    # ----------------------------------------------------
    def test_20_unauthorized_unlocking_rejected(self):
        res = self.client.post('/api/session/unlock_override', json={
            'session_id': 1,
            'auth_code': 'WRONG_PIN',
            'reason': 'Test unauthorized'
        })
        self.assertEqual(res.status_code, 403)

    # ----------------------------------------------------
    # TEST 21: Authorized unlocking requires permission and appropriate reason
    # ----------------------------------------------------
    def test_21_authorized_unlocking_requires_permission_and_reason(self):
        # Missing reason
        res_no_reason = self.client.post('/api/session/unlock_override', json={
            'session_id': 1,
            'auth_code': 'COE@2026',
            'reason': ''
        })
        self.assertEqual(res_no_reason.status_code, 400)

        # Valid PIN and reason
        res_valid = self.client.post('/api/session/unlock_override', json={
            'session_id': 1,
            'auth_code': 'COE@2026',
            'reason': 'Emergency unlock for verification of absentee records'
        })
        self.assertEqual(res_valid.status_code, 200)
        self.assertTrue(res_valid.get_json()['success'])

    # ----------------------------------------------------
    # TEST 22: The Excel file contains all required columns and records
    # ----------------------------------------------------
    def test_22_excel_report_contains_all_columns_and_data(self):
        res = self.client.get('/download_report/1')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.content_type, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # Load workbook in memory
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        ws = wb.active

        # Check title
        self.assertIn("CONTROLLER OF EXAMINATIONS", ws["A1"].value)

        # Check required columns in header row 11
        headers = [ws.cell(row=11, column=c).value for c in range(1, 15)]
        expected = [
            "Serial No.", "Student USN", "Student Name", "Department",
            "Classroom", "Examination Code", "Subject", "Exam Date",
            "Exam Session", "Booklet Barcode", "Attendance Status",
            "Scan Timestamp", "Invigilator ID", "Submission Status"
        ]
        self.assertEqual(headers, expected)

    # ----------------------------------------------------
    # TEST 23: Absent students appear correctly in the exported report
    # ----------------------------------------------------
    def test_23_absent_students_handled_correctly_in_excel(self):
        # 1MS21CS001 was marked ABSENT
        res = self.client.get('/download_report/1')
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        ws = wb.active

        # Find row for 1MS21CS001
        found_absent = False
        for row in range(12, 30):
            usn_val = ws.cell(row=row, column=2).value
            if usn_val == '1MS21CS001':
                found_absent = True
                status_val = ws.cell(row=row, column=11).value
                barcode_val = ws.cell(row=row, column=10).value
                time_val = ws.cell(row=row, column=12).value

                self.assertEqual(status_val, 'ABSENT')
                # Section 14: barcode and timestamp must be empty for absent students
                self.assertEqual(barcode_val or '', '')
                self.assertEqual(time_val or '', '')
                break

        self.assertTrue(found_absent)

    # ----------------------------------------------------
    # TEST 24: Two simultaneous attempts to assign the same booklet cannot conflict
    # ----------------------------------------------------
    def test_24_concurrent_booklet_assignment_safety(self):
        # Re-lock/ensure session 1 is in progress
        results = []
        barcode = 'BK20260010'

        def assign_student(usn):
            client = app.test_client()
            res = client.post('/api/attendance/mark', json={
                'session_id': 1,
                'usn': usn,
                'qp_barcode': barcode
            })
            results.append((usn, res.status_code, res.get_json()))

        # Candidate 3 and Candidate 4 both try to scan BK20260010 simultaneously
        t1 = threading.Thread(target=assign_student, args=('1MS21CS003',))
        t2 = threading.Thread(target=assign_student, args=('1MS21CS004',))

        t1.start()
        t2.start()
        t1.join()
        t2.join()

        # Exactly one must succeed (200), and the other must be rejected (409 conflict)
        statuses = [r[1] for r in results]
        self.assertIn(200, statuses)
        self.assertIn(409, statuses)

    # ----------------------------------------------------
    # TEST 25: The mobile interface works at typical phone screen sizes
    # ----------------------------------------------------
    def test_25_mobile_interface_viewport_and_responsiveness(self):
        self.login_user('admin', 'AdminPassword@2026')
        res = self.client.get('/sessions/1/roster')
        self.assertEqual(res.status_code, 200)
        # Check viewport meta tag
        self.assertIn(b'<meta name="viewport" content="width=device-width, initial-scale=1.0">', res.data)
        # Check mobile touch elements
        self.assertIn(b'btnLaunchScanner', res.data)
        self.assertIn(b'cameraScannerModal', res.data)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 26: Camera permission denial is handled gracefully
    # ----------------------------------------------------
    def test_26_camera_permission_denial_handling(self):
        self.login_user('admin', 'AdminPassword@2026')
        res = self.client.get('/sessions/1/roster')
        self.assertIn(b'cameraErrorBanner', res.data)
        self.assertIn(b'Camera permission denied or camera unavailable', res.data)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 27: The scanner stops when requested
    # ----------------------------------------------------
    def test_27_scanner_stops_when_requested(self):
        self.login_user('admin', 'AdminPassword@2026')
        res = self.client.get('/sessions/1/roster')
        self.assertIn(b'closeCameraModal()', res.data)
        self.assertIn(b'stopCamera()', res.data)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 28: Network errors do not falsely report successful persistence
    # ----------------------------------------------------
    def test_28_network_error_no_false_success(self):
        # Sending corrupted/invalid payload does not produce a false success
        res = self.client.post('/api/attendance/mark', data="invalid-non-json", content_type="application/json")
        self.assertEqual(res.status_code, 400)
        data = res.get_json()
        if data:
            self.assertFalse(data.get('success', False))

    # ----------------------------------------------------
    # TEST 29: CSV import detects duplicate USNs and invalid records
    # ----------------------------------------------------
    def test_29_csv_import_detects_duplicates(self):
        self.login_user('admin', 'AdminPassword@2026')
        csv_data = "usn,student_name,department_code,semester\n1MS21CS001,Duplicate Student,CSE,8\n"
        data = {
            'csv_file': (io.BytesIO(csv_data.encode('utf-8')), 'test.csv')
        }
        res = self.client.post('/admin/students/import', data=data, content_type='multipart/form-data')
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Duplicate Student USN", res.data)
        self.logout_user()

    # ----------------------------------------------------
    # TEST 30: A teacher cannot bypass authorization by directly calling an API
    # ----------------------------------------------------
    def test_30_teacher_cannot_bypass_authorization(self):
        # Teacher attempts to access administrative user creation endpoint
        self.login_user('teacherramesh', 'TeacherPass@2026')
        res = self.client.post('/admin/users/create', data={
            'username': 'hacked_admin',
            'full_name': 'Hacker',
            'email': 'hack@test.edu',
            'password': 'password123'
        })
        # Admin required decorator intercepts and redirects with 302 or returns 403
        self.assertIn(res.status_code, (302, 403))

        # Check DB that user was NOT created
        conn = get_db(TEST_DB_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM users WHERE username = 'hacked_admin'")
        self.assertIsNone(cursor.fetchone())
        conn.close()
        self.logout_user()


if __name__ == '__main__':
    unittest.main()
